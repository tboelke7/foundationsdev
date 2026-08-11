#!/usr/bin/env python3
"""Build Retail_Comparables.xlsx (sortable) and Sources.csv."""
import os, csv
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from common import load_all, group_by_band, band_order, money, show, is_na, ROOT
from content import CATEGORIES, BAND_LABELS
import econ

NAVY = "123A4F"; GOLD = "B8862F"; LIGHT = "F0ECE2"; HEADW = Font(color="FFFFFF", bold=True)
DATE = econ.RESEARCH_DATE
thin = Side(style="thin", color="D8D2C4")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)

CAT_NAME = {c["key"]: c["name"] for c in CATEGORIES}

COLS = [
    ("Category", "category"), ("Band", "band"), ("Product name", "product_name"),
    ("Retailer", "retailer"), ("Price (USD)", "_price"), ("Price shown", "price_current"),
    ("List price", "price_list"), ("Jewelry type", "jewelry_type"),
    ("Gemstone", "gemstone_type"), ("Claimed origin", "claimed_origin"),
    ("Dimensions mm", "stone_dimensions_mm"), ("Carat", "carat_weight"),
    ("Shape / cut", "shape_cut"), ("Treatment", "treatment"), ("Metal", "metal"),
    ("Metal weight", "metal_weight"), ("Setting", "setting_style"),
    ("Construction", "construction"), ("Certification", "certification"),
    ("Positioning", "positioning"), ("Price confidence", "price_confidence"),
    ("Notes", "notes"), ("URL", "url"), ("Image URL", "image_url"),
]
WIDTHS = {"Product name": 40, "Notes": 55, "URL": 45, "Price confidence": 34,
          "Retailer": 20, "Category": 26, "Setting": 20, "Certification": 22,
          "Dimensions mm": 16, "Treatment": 18, "Metal": 18, "Claimed origin": 18,
          "Jewelry type": 18, "Image URL": 30}


def style_header(ws, ncols, row=1):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.font = HEADW
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        cell.border = BORDER


def sheet_comparables(wb, data):
    ws = wb.active
    ws.title = "Comparables"
    ws.append([h for h, _ in COLS])
    style_header(ws, len(COLS))
    band_disp = {**BAND_LABELS}
    for c in CATEGORIES:
        for r in data.get(c["key"], []):
            row = []
            for h, key in COLS:
                if key == "category":
                    row.append(CAT_NAME[c["key"]])
                elif key == "band":
                    row.append(band_disp.get(r.get("_band"), r.get("band", "")))
                elif key == "_price":
                    row.append(r["_price"] if r["_price"] else None)
                else:
                    v = r.get(key)
                    row.append("" if is_na(v) else str(v))
            ws.append(row)
    # formatting
    price_col = [h for h, _ in COLS].index("Price (USD)") + 1
    for rr in range(2, ws.max_row + 1):
        pc = ws.cell(row=rr, column=price_col)
        pc.number_format = '$#,##0'
        for cc in range(1, len(COLS) + 1):
            ws.cell(row=rr, column=cc).alignment = Alignment(vertical="top", wrap_text=True)
            ws.cell(row=rr, column=cc).border = BORDER
    for i, (h, _) in enumerate(COLS, 1):
        ws.column_dimensions[get_column_letter(i)].width = WIDTHS.get(h, 14)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(COLS))}{ws.max_row}"
    return ws.max_row - 1


def sheet_band_analysis(wb, data):
    ws = wb.create_sheet("Band Analysis")
    hdr = ["Category", "Band", "Priced / Total", "Observed prices", "Practical low",
           "Practical high", "Ref retail (mid)", "Ideal outright", "Max outright",
           "Pref consign net", "Max consign net", "Outright profit @ideal",
           "Consign net @pref", "Consign verdict"]
    ws.append(hdr)
    style_header(ws, len(hdr))
    for c in CATEGORIES:
        recs = data.get(c["key"], [])
        for b, rs in group_by_band(c["key"], recs):
            prices = [r["_price"] for r in rs if r["_price"]]
            lo, hi, mid, n = econ.practical_range(prices)
            label = BAND_LABELS.get(b, b)
            if lo and hi:
                R = mid or hi
                o = econ.outright(R); cc = econ.consignment(R)
                ws.append([CAT_NAME[c["key"]], label, f"{n}/{len(rs)}",
                           ", ".join(f"${p:,.0f}" for p in sorted(prices)),
                           lo, hi, R, o["ideal_cost"], o["max_cost"],
                           cc["pref_net"], cc["max_net"], o["profit_at_ideal"],
                           cc["net_profit_pref"], cc["verdict"]])
            else:
                ws.append([CAT_NAME[c["key"]], label, f"0/{len(rs)}",
                           "no price in search", None, None, None, None, None,
                           None, None, None, None, "no comps"])
    for rr in range(2, ws.max_row + 1):
        for col in (5, 6, 7, 8, 9, 10, 11, 12, 13):
            ws.cell(row=rr, column=col).number_format = '$#,##0'
        for cc in range(1, len(hdr) + 1):
            ws.cell(row=rr, column=cc).border = BORDER
            ws.cell(row=rr, column=cc).alignment = Alignment(vertical="top", wrap_text=True)
    widths = [26, 22, 12, 40, 12, 12, 13, 13, 13, 14, 14, 16, 15, 13]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"


def sheet_quickref(wb):
    ws = wb.create_sheet("Quick Reference")
    hdr = ["Expected U.S. Retail", "Ideal Outright Cost", "Max Outright Cost",
           "Preferred Consignment Net", "Max Consignment Net", "Landed Allowance",
           "Outright Profit @ideal", "Outright Margin @ideal",
           "Consign Net Profit @pref", "Consign Margin @pref"]
    ws.append(hdr)
    style_header(ws, len(hdr))
    for r in econ.quick_ref_rows():
        ws.append([r["retail"], r["ideal_outright"], r["max_outright"],
                   r["pref_consign_net"], r["max_consign_net"], r["landed_allowance"],
                   r["profit_outright_ideal"], r["margin_outright_ideal"],
                   r["net_consign_pref"], r["margin_consign_pref"]])
    for rr in range(2, ws.max_row + 1):
        for col in range(1, 8):
            ws.cell(row=rr, column=col).number_format = '$#,##0'
        for col in (8, 10):
            ws.cell(row=rr, column=col).number_format = '0%'
        for cc in range(1, len(hdr) + 1):
            ws.cell(row=rr, column=cc).border = BORDER
    for i, w in enumerate([20, 18, 18, 22, 20, 16, 18, 18, 20, 18], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"


def sheet_assumptions(wb):
    ws = wb.create_sheet("Assumptions")
    ws.append(["Assumption", "Value", "Notes"])
    style_header(ws, 3)
    A = econ.A
    rows = [
        ("Ideal outright cost", A["outright_ideal_pct"], "fraction of expected U.S. retail (R)"),
        ("Max outright cost", A["outright_max_pct"], "fraction of R — do not exceed"),
        ("Preferred consignment net", A["consign_pref_pct"], "fraction of R paid to SL jeweler"),
        ("Max consignment net", A["consign_max_pct"], "fraction of R — upper warning level"),
        ("Payment processing", A["payment_processing_pct"], "fraction of R"),
        ("Returns / damage reserve", A["returns_reserve_pct"], "fraction of R"),
        ("Insured intl shipping floor", A["intl_ship_floor"], "USD minimum per insured parcel SL->US"),
        ("Insured intl shipping %", A["intl_ship_pct"], "added fraction of R for higher value"),
        ("Import duty allowance", A["customs_duty_pct"], "fraction of value; finished jewelry HTS 7113 ~5-5.8%; loose stones often 0%"),
        ("Packaging", A["packaging_flat"], "USD per piece"),
        ("Domestic shipping", A["domestic_ship_flat"], "USD per U.S. sale"),
    ]
    for k, v, note in rows:
        ws.append([k, v, note])
    for rr in range(2, ws.max_row + 1):
        val = ws.cell(row=rr, column=2).value
        ws.cell(row=rr, column=2).number_format = '0%' if isinstance(val, float) and val < 1 else '$#,##0'
    ws.append([])
    ws.append(["Outright formula", "profit = R - acquisition - (intl_ship + duty*R + packaging) - (3%R + 5%R + $12)"])
    ws.append(["Consignment formula", "net = R - jeweler_net - (3%R + intl_ship + duty*R + packaging + 5%R)"])
    ws.append(["Research date", DATE])
    ws.append(["IMPORTANT", "All prices are ASKING prices observed via web search, not confirmed sale values. Verify each at its URL."])
    ws.append(["Small pieces", "Under ~$300 retail the fixed ~$60 insured-shipping floor eats the margin — only import commodity pieces in BULK to amortize freight."])
    for i, w in enumerate([32, 16, 80], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.column_dimensions["A"].width = 30


def build_sources_csv(data):
    path = os.path.join(ROOT, "Sources.csv")
    with open(path, "w", newline="") as f:
        w = csv.writer(f)
        w.writerow(["Category", "Band", "Product name", "Retailer", "Positioning",
                    "Price shown (USD)", "URL", "Research date", "Price confidence"])
        for c in CATEGORIES:
            for r in data.get(c["key"], []):
                w.writerow([CAT_NAME[c["key"]], BAND_LABELS.get(r.get("_band"), r.get("band", "")),
                            show(r.get("product_name")), show(r.get("retailer")),
                            show(r.get("positioning")),
                            (f"{r['_price']:.2f}" if r["_price"] else ""),
                            show(r.get("url")), DATE, show(r.get("price_confidence"))])
    return path


def main():
    data = load_all()
    wb = Workbook()
    n = sheet_comparables(wb, data)
    sheet_band_analysis(wb, data)
    sheet_quickref(wb)
    sheet_assumptions(wb)
    # Sources sheet inside workbook too
    ws = wb.create_sheet("Sources")
    ws.append(["Category", "Product", "Retailer", "Price (USD)", "URL", "Researched"])
    style_header(ws, 6)
    for c in CATEGORIES:
        for r in data.get(c["key"], []):
            ws.append([CAT_NAME[c["key"]], show(r.get("product_name")), show(r.get("retailer")),
                       (r["_price"] if r["_price"] else None), show(r.get("url")), DATE])
    for rr in range(2, ws.max_row + 1):
        ws.cell(row=rr, column=4).number_format = '$#,##0'
    for i, w in enumerate([26, 42, 20, 12, 48, 12], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    xlsx = os.path.join(ROOT, "Retail_Comparables.xlsx")
    wb.save(xlsx)
    csvp = build_sources_csv(data)
    print(f"Wrote {xlsx} ({n} comparable rows) and {csvp}")


if __name__ == "__main__":
    main()
